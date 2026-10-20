#!/usr/bin/env python3
"""Data Cloud inventory generator.

Pulls every Data Cloud data stream in an org and enriches it with its DLO
(Data Lake Object) and mapped DMO(s), then writes two filterable artifacts:

  * data_stream_inventory.xlsx  -- multi-sheet workbook, each sheet with a
    frozen header row and column autofilter dropdowns.
  * data_stream_inventory.html  -- one self-contained page with live search,
    per-column filters, and sortable columns (no dependencies).

Auth is borrowed from the Salesforce CLI (`sf`); no secrets live in this file.
Run it yourself anytime:

    python generate_inventory.py --org jdo-oe0sdd

See README.md for the column guide and method caveats.
"""
from __future__ import annotations

import argparse
import os
import concurrent.futures as cf
import html
import json
import re
import subprocess
import sys
import threading
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict

# ---------------------------------------------------------------------------
# Salesforce CLI auth
# ---------------------------------------------------------------------------

_TOKEN_RE = re.compile(r"00D[A-Za-z0-9!._-]{20,}")


def _run(cmd: list[str]) -> str:
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        raise SystemExit(
            f"command failed ({proc.returncode}): {' '.join(cmd)}\n{proc.stderr.strip()}"
        )
    return proc.stdout


def get_session(org: str) -> tuple[str, str]:
    """Return (instance_url, access_token) for the org alias via the sf CLI."""
    disp = json.loads(_run(["sf", "org", "display", "-o", org, "--json"]))
    instance = disp["result"]["instanceUrl"].rstrip("/")
    # `sf org display` redacts the token on modern CLIs; ask for it explicitly.
    raw = _run(["sf", "org", "auth", "show-access-token", "-o", org, "--no-prompt"])
    m = _TOKEN_RE.search(raw)
    if not m:
        raise SystemExit(
            "could not parse an access token from `sf org auth show-access-token`.\n"
            f"raw output:\n{raw}"
        )
    return instance, m.group(0)


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------


class Client:
    def __init__(self, instance: str, token: str, api: str):
        self.base = f"{instance}/services/data/v{api}"
        self.headers = {"Authorization": f"Bearer {token}"}

    def get(self, path: str, params: dict | None = None, retries: int = 2,
            timeout: int = 25):
        url = self.base + path
        if params:
            url += "?" + urllib.parse.urlencode(params)
        last = None
        for attempt in range(retries):
            req = urllib.request.Request(url, headers=self.headers)
            try:
                with urllib.request.urlopen(req, timeout=timeout) as resp:
                    body = resp.read().decode().strip()
                    return json.loads(body) if body else {}  # empty 200 == no results
            except urllib.error.HTTPError as e:
                body = e.read().decode(errors="replace")
                last = f"HTTP {e.code}: {body[:300]}"
                if e.code in (401, 403):
                    raise SystemExit(f"{url}\n{last}")
                if e.code in (400, 404):  # bad/unknown lookup key -> no results
                    return {}
            except Exception as e:  # noqa: BLE001 - network flakiness, retry once
                last = str(e)
        raise SystemExit(f"GET failed after {retries} tries: {url}\n{last}")


# ---------------------------------------------------------------------------
# Fetchers
# ---------------------------------------------------------------------------


def fetch_streams(client: Client, page: int = 50) -> list[dict]:
    out: list[dict] = []
    offset = 1
    while True:
        # Each stream embeds its full DLO + source field lists, so this payload
        # is heavy -- keep pages small and give the call extra time.
        d = client.get("/ssot/data-streams", {"limit": page, "offset": offset},
                       retries=3, timeout=90)
        batch = d.get("dataStreams") or []
        out.extend(batch)
        total = d.get("totalSize", len(out))
        print(f"  streams {len(out)}/{total}", file=sys.stderr)
        if len(out) >= total or not batch:
            break
        offset += len(batch)
    return out


def fetch_dmos(client: Client) -> dict[str, dict]:
    """name -> {category, fieldCount, primaryKeys, displayName}."""
    d = client.get("/ssot/metadata", {"entityType": "DataModelObject"},
                   retries=3, timeout=120)
    out: dict[str, dict] = {}
    for it in d.get("metadata") or []:
        name = it.get("name")
        if not name:
            continue
        out[name] = {
            "category": it.get("category"),
            "displayName": it.get("displayName"),
            "fieldCount": len(it.get("fields") or []),
            "primaryKeys": [p.get("name") if isinstance(p, dict) else p
                            for p in (it.get("primaryKeys") or [])],
        }
    return out


def fetch_dlos(client: Client) -> dict[str, dict]:
    """All Data Lake Objects, so we can flag orphans. Best-effort.

    name -> {category, status, fieldCount}. Returns {} if the entity type is
    unavailable on this org/API version.
    """
    try:
        d = client.get("/ssot/metadata", {"entityType": "DataLakeObject"},
                       retries=3, timeout=120)
    except SystemExit as e:
        print(f"  (DLO list unavailable, skipping orphan detection: {e})", file=sys.stderr)
        return {}
    out: dict[str, dict] = {}
    for it in d.get("metadata") or []:
        name = it.get("name")
        if not name:
            continue
        out[name] = {
            "category": it.get("category"),
            "status": it.get("status"),
            "fieldCount": len(it.get("fields") or []),
        }
    return out


def fetch_mappings(client: Client, dmos: dict[str, dict], workers: int) -> list[dict]:
    """Resolve DLO->DMO mappings by querying each DMO (reverse lookup).

    We enumerate by DMO rather than by source object on purpose: the
    source-object (forward) query on large CRM objects returns multi-megabyte
    payloads that time out / reset the connection, whereas the per-DMO query
    returns only that DMO's maps and is fast. Reverse still covers CRM streams
    (e.g. dmoDeveloperName=ssot__Individual__dlm surfaces Contact_Home_map_…).
    Maps are deduped by their developerName.
    """
    maps: dict[str, dict] = {}
    lock = threading.Lock()

    def pull(name: str):
        d = client.get("/ssot/data-model-object-mappings", {"dmoDeveloperName": name})
        rows = d.get("objectSourceTargetMaps") if isinstance(d, dict) else None
        if not rows:
            return
        with lock:
            for r in rows:
                dev = r.get("developerName")
                if dev:
                    maps[dev] = r

    names = list(dmos)
    done = 0
    with cf.ThreadPoolExecutor(max_workers=workers) as ex:
        futs = [ex.submit(pull, n) for n in names]
        for _ in cf.as_completed(futs):
            done += 1
            if done % 100 == 0 or done == len(futs):
                print(f"  mappings {done}/{len(futs)} DMOs queried", file=sys.stderr)
    return list(maps.values())


# ---------------------------------------------------------------------------
# Join
# ---------------------------------------------------------------------------


def map_prefix(dev_name: str) -> str | None:
    """`Account_Home_map_Account_1736..` -> `Account_Home` (the stream/DLO base)."""
    if "_map_" not in dev_name:
        return None
    return dev_name.rsplit("_map_", 1)[0]


def build_stream_dmos(streams: list[dict], maps: list[dict]) -> dict[str, set[str]]:
    """stream name -> set of mapped DMO developer names."""
    # Match a map's prefix to a stream via the stream name or its DLO base name.
    by_key: dict[str, str] = {}
    for s in streams:
        name = s.get("name")
        by_key[name] = name
        dlo = (s.get("dataLakeObjectInfo") or {}).get("name")  # e.g. Foo__dll
        if dlo:
            by_key[dlo] = name
            by_key[dlo.removesuffix("__dll")] = name

    result: dict[str, set[str]] = defaultdict(set)
    for m in maps:
        dmo = m.get("targetEntityDeveloperName")
        prefix = map_prefix(m.get("developerName") or "")
        if not dmo or not prefix:
            continue
        stream = by_key.get(prefix)
        if stream is None:
            # longest-prefix fallback (handles rare suffix drift)
            cand = [k for k in by_key if prefix.startswith(k)]
            if cand:
                stream = by_key[max(cand, key=len)]
        if stream is not None:
            result[stream].add(dmo)
    return result


# ---------------------------------------------------------------------------
# Persona tagging
# ---------------------------------------------------------------------------

DEFAULT_RULES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "persona_rules.json")


class PersonaTagger:
    """Tags each stream with the persona(s) it serves, per persona_rules.json.

    Exact `overrides` win; otherwise the first regex rule matching the stream
    name, label, or source object applies; otherwise `Unclassified`.

    `suppress_rules` flag streams that aren't customer-useful (test, internal,
    initiative-specific). Pattern rules are checked first, then persona rules
    (`personas`), which run against the final -- possibly corrected -- tags.
    """

    def __init__(self, path: str):
        with open(path, encoding="utf-8") as f:
            cfg = json.load(f)
        self.personas: list[str] = cfg.get("personas") or []
        self.overrides: dict[str, list[str]] = cfg.get("overrides") or {}
        self.rules = [(r["tags"], r.get("label") or r["pattern"],
                       re.compile(r["pattern"], re.IGNORECASE))
                      for r in cfg.get("rules") or []]
        sup = cfg.get("suppress_rules") or []
        self.suppress_patterns = [(r["reason"], re.compile(r["pattern"], re.IGNORECASE))
                                  for r in sup if r.get("pattern")]
        self.suppress_personas = [(r["personas"], r["reason"]) for r in sup if r.get("personas")]
        self.tag_choices = sorted({t for tags, _, _ in self.rules for t in tags}
                                  | {t for v in self.overrides.values() for t in v})

    @staticmethod
    def _fields(stream: dict) -> list[str]:
        src = ((stream.get("connectorInfo") or {}).get("connectorDetails") or {}).get("sourceObject")
        return [f for f in (stream.get("name"), stream.get("label"), src) if f]

    def tag(self, stream: dict) -> tuple[list[str], str]:
        name = stream.get("name") or ""
        if name in self.overrides:
            return list(self.overrides[name]), "override"
        fields = self._fields(stream)
        for tags, label, rx in self.rules:
            if any(rx.search(f) for f in fields):
                return list(tags), f"rule: {label}"
        return ["Unclassified"], "no rule matched"

    def pattern_suppression(self, stream: dict) -> str | None:
        fields = self._fields(stream)
        for reason, rx in self.suppress_patterns:
            if any(rx.search(f) for f in fields):
                return reason
        return None

    def persona_suppression(self, tags: list[str]) -> str | None:
        for personas, reason in self.suppress_personas:
            if set(personas) & set(tags):
                return reason
        return None


def load_corrections(path: str | None) -> dict:
    """Collaborator corrections (exported from the HTML editor), or empty."""
    if not path or not os.path.exists(path):
        return {"streams": {}}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    data.setdefault("streams", {})
    print(f"  {len(data['streams'])} corrections from {path}", file=sys.stderr)
    return data


ATTENTION_RUN = {"FAILURE", "FAILED", "WARNING", "ERROR", "PARTIAL_SUCCESS"}


def build_rows(streams: list[dict], stream_dmos: dict[str, set[str]],
               dmos: dict[str, dict], tagger: PersonaTagger,
               corrections: dict[str, dict]) -> tuple[list[dict], dict[str, dict]]:
    """Returns (rows, rule_values) -- the latter is the rules-only verdict per
    stream, which the HTML editor needs for "revert to rules"."""
    rows, rule_values = [], {}
    for s in streams:
        conn = (s.get("connectorInfo") or {})
        cd = conn.get("connectorDetails", {}) or {}
        dlo = s.get("dataLakeObjectInfo") or {}
        ds = dlo.get("dataSpaceInfo")
        if isinstance(ds, list):
            spaces = "; ".join(sorted({(d.get("name") or d.get("label")) for d in ds
                                       if isinstance(d, dict)} - {None}))
        elif isinstance(ds, dict):
            spaces = ds.get("name") or ds.get("label")
        else:
            spaces = None
        mapped = sorted(stream_dmos.get(s.get("name"), set()))
        # field count across mapped DMOs (max, they usually share the target)
        dmo_fields = max((dmos.get(d, {}).get("fieldCount", 0) for d in mapped),
                         default=None)
        run = (s.get("lastRunStatus") or "").upper()
        status = (s.get("status") or "").upper()
        enabled = s.get("isEnabled")
        needs_attn = (run in ATTENTION_RUN) or (status and status != "ACTIVE") or (enabled is False)
        name = s.get("name")
        rule_tags, rule_basis = tagger.tag(s)
        pattern_reason = tagger.pattern_suppression(s)
        rule_reason = pattern_reason or tagger.persona_suppression(rule_tags)
        rule_values[name] = {"p": "; ".join(rule_tags), "b": rule_basis,
                             "s": bool(rule_reason), "r": rule_reason or "",
                             "pr": pattern_reason}
        # A correction may override persona and/or suppression independently;
        # mirrored in the HTML's applyRow().
        c = corrections.get(name) or {}
        personas = c.get("persona") or rule_tags
        basis = "correction" if c.get("persona") else rule_basis
        if "suppressed" in c:
            suppressed, reason = bool(c["suppressed"]), (c.get("reason") or "")
        else:
            reason = pattern_reason or tagger.persona_suppression(personas) or ""
            suppressed = bool(reason)
        rows.append({
            "Data Stream Name": s.get("name"),
            "Data Stream Label": s.get("label"),
            "Persona": "; ".join(personas),
            "Persona Basis": basis,
            "Suppressed": suppressed,
            "Suppress Reason": reason,
            "Notes": c.get("note") or "",
            "Corrected By": c.get("by") or "",
            "Stream Status": s.get("status"),
            "Enabled": enabled,
            "Last Run Status": s.get("lastRunStatus"),
            "Last Refresh Date": s.get("lastRefreshDate"),
            "Total Records": s.get("totalRecords"),
            "Last Added Records": s.get("lastAddedRecords"),
            "Connector Type": conn.get("connectorType") or cd.get("type"),
            "Connector Name": cd.get("name"),
            "Data Source": s.get("dataSource"),
            "Data Stream Type": s.get("dataStreamType"),
            "Source Object": cd.get("sourceObject"),
            "Refresh Mode": (s.get("refreshConfig") or {}).get("refreshMode"),
            "Frequency Type": ((s.get("refreshConfig") or {}).get("frequency") or {}).get("frequencyType"),
            "DLO Name": dlo.get("name"),
            "DLO Category": dlo.get("category"),
            "DLO Status": dlo.get("status"),
            "DLO Field Count": len(dlo.get("dataLakeFieldInfoRepresentation") or []),
            "Data Space": spaces,
            "Mapped DMO(s)": "; ".join(mapped),
            "Mapping Count": len(mapped),
            "DMO Field Count": dmo_fields,
            "Is Mapped": bool(mapped),
            "Needs Attention": needs_attn,
            "Record ID": s.get("recordId"),
        })
    rows.sort(key=lambda r: (r["Data Stream Name"] or "").lower())
    return rows, rule_values


# ---------------------------------------------------------------------------
# Breakdowns
# ---------------------------------------------------------------------------


def build_breakdowns(rows: list[dict], streams: list[dict], dlos: dict[str, dict],
                     stream_dmos: dict[str, set[str]]):
    unmapped = [r for r in rows if not r["Is Mapped"]]
    attention = [r for r in rows if r["Needs Attention"]]

    # Orphaned DLOs: exist in the org but no visible stream feeds them.
    orphaned = []
    if dlos:
        stream_dlo_names = {(s.get("dataLakeObjectInfo") or {}).get("name") for s in streams}
        for name, meta in sorted(dlos.items()):
            if name not in stream_dlo_names:
                orphaned.append({
                    "DLO Name": name,
                    "DLO Category": meta.get("category"),
                    "DLO Status": meta.get("status"),
                    "Field Count": meta.get("fieldCount"),
                })

    # Summaries.
    def summarize(key: str):
        agg = defaultdict(lambda: {"Streams": 0, "Mapped": 0, "Needs Attention": 0,
                                   "Total Records": 0})
        for r in rows:
            k = r.get(key) or "(none)"
            a = agg[k]
            a["Streams"] += 1
            a["Mapped"] += 1 if r["Is Mapped"] else 0
            a["Needs Attention"] += 1 if r["Needs Attention"] else 0
            a["Total Records"] += r.get("Total Records") or 0
        return [{key: k, **v} for k, v in sorted(agg.items(),
                key=lambda kv: -kv[1]["Streams"])]

    return {
        "Suppressed Streams": [r for r in rows if r["Suppressed"]],
        "Unmapped Streams": unmapped,
        "Needs Attention": attention,
        "Orphaned DLOs": orphaned,
        "Summary by Persona": summarize_personas(rows),
        "Summary by Connector": summarize("Connector Type"),
        "Summary by Category": summarize("DLO Category"),
        "Summary by Source Object": summarize("Source Object"),
    }


def persona_tags(row: dict) -> list[str]:
    return [t.strip() for t in (row.get("Persona") or "").split(";") if t.strip()]


def persona_rows(rows: list[dict], persona: str) -> list[dict]:
    """Streams that apply to `persona` -- its own tag plus shared `All` data."""
    return [r for r in rows if {persona, "All"} & set(persona_tags(r))]


def summarize_personas(rows: list[dict]) -> list[dict]:
    """One line per tag; a multi-tagged stream counts toward each of its tags."""
    agg = defaultdict(lambda: {"Streams": 0, "Suppressed": 0, "Mapped": 0,
                               "Needs Attention": 0, "Total Records": 0})
    for r in rows:
        for t in persona_tags(r):
            a = agg[t]
            a["Streams"] += 1
            a["Suppressed"] += 1 if r["Suppressed"] else 0
            a["Mapped"] += 1 if r["Is Mapped"] else 0
            a["Needs Attention"] += 1 if r["Needs Attention"] else 0
            a["Total Records"] += r.get("Total Records") or 0
    return [{"Persona": k, **v} for k, v in sorted(agg.items(), key=lambda kv: -kv[1]["Streams"])]


def build_connectors(streams: list[dict]) -> list[dict]:
    agg: dict[str, dict] = {}
    for s in streams:
        conn = s.get("connectorInfo") or {}
        cd = conn.get("connectorDetails", {}) or {}
        name = cd.get("name") or "(none)"
        a = agg.setdefault(name, {"Connector Name": name,
                                  "Connector Type": conn.get("connectorType") or cd.get("type"),
                                  "Stream Count": 0, "_objs": set()})
        a["Stream Count"] += 1
        if cd.get("sourceObject"):
            a["_objs"].add(cd["sourceObject"])
    out = []
    for a in sorted(agg.values(), key=lambda x: -x["Stream Count"]):
        objs = sorted(a.pop("_objs"))
        a["Source Object Count"] = len(objs)
        a["Source Objects"] = "; ".join(objs)
        out.append(a)
    return out


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def write_xlsx(path: str, sheets: dict[str, list[dict]]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")

    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        if not rows:
            ws.append(["(no rows)"])
            continue
        cols = list(rows[0].keys())
        ws.append(cols)
        for r in rows:
            ws.append([_cell(r.get(c)) for c in cols])
        for ci, _ in enumerate(cols, 1):
            c = ws.cell(row=1, column=ci)
            c.font = head_font
            c.fill = head_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        for ci, col in enumerate(cols, 1):
            width = max(len(str(col)),
                        *(len(str(_cell(r.get(col)))) for r in rows[:200])) + 2
            ws.column_dimensions[get_column_letter(ci)].width = min(width, 55)
    wb.save(path)


def _cell(v):
    if isinstance(v, (list, set)):
        return "; ".join(map(str, v))
    return v


def write_html(path: str, title: str, rows: list[dict], org: str, generated: str,
               tagger: PersonaTagger, corrections: dict, rule_values: dict):
    cols = list(rows[0].keys()) if rows else []
    data_json = json.dumps([[_cell(r.get(c)) for c in cols] for r in rows])
    cols_json = json.dumps(cols)
    tmpl = _HTML.replace("__TITLE__", html.escape(title))
    tmpl = tmpl.replace("__ORG__", html.escape(org))
    tmpl = tmpl.replace("__GEN__", html.escape(generated))
    tmpl = tmpl.replace("__COUNT__", str(len(rows)))
    tmpl = tmpl.replace("__COLS__", cols_json)
    tmpl = tmpl.replace("__PERSONAS__", json.dumps(tagger.personas))
    tmpl = tmpl.replace("__TAGS__", json.dumps(tagger.tag_choices))
    tmpl = tmpl.replace("__SUPP_P__", json.dumps(tagger.suppress_personas))
    tmpl = tmpl.replace("__ORGJS__", json.dumps(org))
    tmpl = tmpl.replace("__CORR__", json.dumps(corrections))
    tmpl = tmpl.replace("__RULE__", json.dumps(rule_values))
    tmpl = tmpl.replace("__DATA__", data_json)
    with open(path, "w", encoding="utf-8") as f:
        f.write(tmpl)


_HTML = r"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title>
<style>
  :root{--bd:#d0d7de;--hd:#1F4E78;--zebra:#f6f8fa;--accent:#0969da}
  *{box-sizing:border-box}
  body{font:13px/1.4 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;margin:0;color:#1f2328;background:#fff}
  header{padding:14px 18px;border-bottom:1px solid var(--bd);position:sticky;top:0;background:#fff;z-index:5}
  h1{font-size:16px;margin:0 0 4px}
  .meta{color:#57606a;font-size:12px}
  .bar{margin-top:10px;display:flex;gap:10px;align-items:center;flex-wrap:wrap}
  select{padding:7px 8px;border:1px solid var(--bd);border-radius:6px;font-size:12px;background:#fff}
  #q{flex:1;min-width:220px;padding:7px 10px;border:1px solid var(--bd);border-radius:6px;font-size:13px}
  button{padding:7px 12px;border:1px solid var(--bd);background:#f6f8fa;border-radius:6px;cursor:pointer;font-size:12px}
  #count{color:#57606a;font-size:12px}
  .wrap{overflow:auto;max-height:calc(100vh - 120px)}
  table{border-collapse:collapse;width:100%}
  th,td{border:1px solid var(--bd);padding:5px 8px;text-align:left;white-space:nowrap;max-width:420px;overflow:hidden;text-overflow:ellipsis}
  thead th{position:sticky;top:0;background:var(--hd);color:#fff;cursor:pointer;z-index:2}
  thead tr.filters th{position:sticky;top:29px;background:#eef2f6;z-index:1}
  thead tr.filters input{width:100%;box-sizing:border-box;border:1px solid var(--bd);border-radius:4px;padding:3px 5px;font-size:11px}
  tbody tr:nth-child(even){background:var(--zebra)}
  tbody tr:hover{background:#fff8c5}
  th .arrow{font-size:10px;opacity:.7;margin-left:4px}
  td.bool-true{color:#1a7f37;font-weight:600}
  td.bool-false{color:#cf222e}
  body.editing tbody tr{cursor:pointer}
  body.editing #edit{background:#0969da;color:#fff;border-color:#0969da}
  tr.pending td{background:#ddf4ff!important}
  #pend{color:#0969da;font-weight:600;font-size:12px}
  dialog{border:1px solid var(--bd);border-radius:8px;padding:16px 18px;width:440px;max-width:95vw}
  dialog h3{margin:0 0 2px;font-size:14px;word-break:break-all}
  dialog .sub{color:#57606a;font-size:12px;margin-bottom:10px}
  dialog fieldset{border:1px solid var(--bd);border-radius:6px;margin:0 0 10px;padding:8px 10px}
  dialog fieldset label{display:inline-block;margin:2px 12px 2px 0}
  dialog input[type=text],dialog textarea,dialog select{width:100%;box-sizing:border-box;margin-top:4px;padding:5px 7px;border:1px solid var(--bd);border-radius:5px;font:inherit}
  body.editing td.ed{background:#fffbea;cursor:text}
  body.editing td.tog{cursor:pointer;text-decoration:underline dotted}
  td.editing{outline:2px solid var(--accent);white-space:normal;background:#fff!important}
  th.sel,td.sel{width:28px;text-align:center;padding:0 4px}
  tr.selected td{background:#fff1b8!important}
  #selbar{font-size:12px;display:inline-flex;gap:6px;align-items:center}
  dialog .row{display:flex;gap:8px;margin-top:12px;justify-content:flex-end}
</style></head><body>
<header>
  <h1>__TITLE__</h1>
  <div class="meta">Org <b>__ORG__</b> &middot; generated __GEN__ &middot; <span id="count">__COUNT__</span> rows</div>
  <div class="bar">
    <input id="q" placeholder="Search all columns&hellip;" autofocus>
    <select id="persona" title="Show streams for a persona (includes shared 'All' data)"></select>
    <button id="cols">Toggle filters</button>
    <button id="reset">Reset</button>
    <button id="csv" title="Export the rows currently shown">Export CSV</button>
    <button id="xlsx" title="Export the rows currently shown">Export Excel</button>
    <button id="edit" title="Click any row to correct its persona or suppression">Edit classifications</button>
    <span id="selbar" hidden><b id="selN"></b> selected
      <button id="bulk">Bulk edit&hellip;</button><button id="selClear">Clear selection</button></span>
    <span id="pend"></span>
    <button id="exp" hidden title="Download the merged corrections file to commit to git">Export corrections</button>
    <button id="disc" hidden>Discard pending</button>
  </div>
</header>
<dialog id="ed"><form method="dialog">
  <h3 id="edName"></h3><div class="sub" id="edRule"></div>
  <fieldset><legend>Persona</legend><div id="edP"></div></fieldset>
  <fieldset><legend>Suppression</legend>
    <label><input type="checkbox" id="edS"> Suppressed &mdash; not customer-useful</label>
    <select id="edR" class="reasons"></select>
    <input type="text" id="edRx" placeholder="detail, e.g. React cockpits">
  </fieldset>
  <textarea id="edN" rows="2" placeholder="Note for collaborators (optional)"></textarea>
  <div class="row"><button value="revert">Revert to rules</button><button value="cancel">Cancel</button><button value="save">Save</button></div>
</form></dialog>
<dialog id="bk"><form method="dialog">
  <h3 id="bkTitle"></h3><div class="sub">Only the sections you change are applied; everything else is left as-is per row.</div>
  <fieldset><legend>Persona</legend>
    <select id="bkPm"><option value="">Leave unchanged</option><option value="set">Replace with</option>
      <option value="add">Add</option><option value="remove">Remove</option></select>
    <div id="bkP"></div></fieldset>
  <fieldset><legend>Suppression</legend>
    <select id="bkSm"><option value="">Leave unchanged</option><option value="on">Suppress</option><option value="off">Unsuppress</option></select>
    <select id="bkR" class="reasons"></select>
    <input type="text" id="bkRx" placeholder="detail, e.g. React cockpits"></fieldset>
  <fieldset><legend>Note</legend>
    <select id="bkNm"><option value="">Leave unchanged</option><option value="set">Set to</option><option value="clear">Clear</option></select>
    <input type="text" id="bkN" placeholder="note"></fieldset>
  <div class="row"><button value="revert">Revert all to rules</button><button value="cancel">Cancel</button><button value="apply" id="bkGo">Apply</button></div>
</form></dialog>
<div class="wrap"><table id="t"><thead></thead><tbody></tbody></table></div>
<script>
const COLS=__COLS__, DATA=__DATA__, PERSONAS=__PERSONAS__;
const PCOL=COLS.indexOf("Persona");
// --- Collaborative corrections: CORR = committed file, pending = this browser's edits ---
const TAGS=__TAGS__, SUPP_P=__SUPP_P__, ORG=__ORGJS__, CORR=__CORR__, RULE=__RULE__;
const C=n=>COLS.indexOf(n), NAME=C("Data Stream Name"), PB=C("Persona Basis"),
  SUP=C("Suppressed"), SR=C("Suppress Reason"), NOTE=C("Notes"), BY=C("Corrected By");
const LSK="dcinv-corrections-"+ORG;
const same=(a,b)=>JSON.stringify(a??null)===JSON.stringify(b??null);
const strip=e=>{if(!e)return null;const {by,at,...x}=e;return x;};
const ruleReasonFor=(r,tags)=>{const pr=SUPP_P.find(([ps])=>ps.some(p=>tags.includes(p)));return r.pr||(pr?pr[1]:"");};
const sameTags=(a,b)=>[...a].sort().join("; ")===[...b].sort().join("; ");
// Drop the parts of a correction the current rules already give (null = nothing left).
function minimal(n,e){const r=RULE[n];if(!e||!r)return e;const x={...e};
  if(x.persona&&(!x.persona.length||sameTags(x.persona,r.p.split("; "))))delete x.persona;
  if("suppressed" in x){const rr=ruleReasonFor(r,x.persona||r.p.split("; "));
    if(!!x.suppressed===!!rr&&(!x.suppressed||x.reason===rr)){delete x.suppressed;delete x.reason;}}
  if(!x.note)delete x.note;
  return Object.keys(strip(x)).length?x:null;}
let pending=JSON.parse(localStorage.getItem(LSK)||"{}");
// Saved edits go stale after a commit or rule change: shrink each to what the rules
// don't already give, and drop it if that matches the committed file.
for(const n of Object.keys(pending)){
  const m=pending[n]===null?null:minimal(n,pending[n]);
  if(same(strip(m),strip(minimal(n,CORR.streams[n]))))delete pending[n];else pending[n]=m;}
localStorage.setItem(LSK,JSON.stringify(pending));
const corrFor=n=>(n in pending)?pending[n]:CORR.streams[n];
function applyRow(row){  // mirrors build_rows() in generate_inventory.py
  const n=row[NAME],r=RULE[n],c=corrFor(n)||{};
  const tags=c.persona&&c.persona.length?c.persona:r.p.split("; ");
  row[PCOL]=tags.join("; ");row[PB]=c.persona&&c.persona.length?"correction":r.b;
  if("suppressed" in c){row[SUP]=!!c.suppressed;row[SR]=c.reason||"";}
  else{const pr=SUPP_P.find(([ps])=>ps.some(p=>tags.includes(p)));
    const reason=r.pr||(pr?pr[1]:"");row[SUP]=!!reason;row[SR]=reason;}
  row[NOTE]=c.note||"";row[BY]=c.by||"";
}
DATA.forEach(applyRow);
const tagsOf=row=>String(row[PCOL]||"").split(";").map(t=>t.trim()).filter(Boolean);
const personaSel=document.getElementById("persona");
(()=>{const all=new Set();DATA.forEach(r=>tagsOf(r).forEach(t=>all.add(t)));
  const opts=[["","All personas"],...PERSONAS.map(p=>[p,p+" (incl. All)"]),
    ...[...all].filter(t=>!PERSONAS.includes(t)).sort().map(t=>["="+t,"Only: "+t])];
  personaSel.innerHTML=opts.map(([v,l])=>`<option value="${v}">${l}</option>`).join("");})();
let sortCol=-1, sortDir=1, showFilters=true;
const colFilters=COLS.map(()=>"");
const thead=document.querySelector("thead"), tbody=document.querySelector("tbody");
const q=document.getElementById("q"), count=document.getElementById("count");
function buildHead(){
  thead.innerHTML="";
  const h=document.createElement("tr");
  if(editing){const th=document.createElement("th");th.className="sel";
    th.innerHTML='<input type="checkbox" id="selAll" title="Select all rows shown">';
    th.onclick=e=>e.stopPropagation();h.appendChild(th);}
  COLS.forEach((c,i)=>{const th=document.createElement("th");
    th.innerHTML=c+'<span class="arrow"></span>';
    th.onclick=()=>{sortDir=(sortCol===i)?-sortDir:1;sortCol=i;render();};
    h.appendChild(th);});
  thead.appendChild(h);
  const f=document.createElement("tr");f.className="filters";
  if(editing)f.appendChild(document.createElement("th"));
  COLS.forEach((c,i)=>{const th=document.createElement("th");
    const inp=document.createElement("input");inp.placeholder="filter";
    inp.value=colFilters[i];
    inp.oninput=e=>{colFilters[i]=e.target.value.toLowerCase();render();};
    th.appendChild(inp);f.appendChild(th);});
  f.style.display=showFilters?"":"none";
  thead.appendChild(f);
}
function filtered(){
  const g=q.value.toLowerCase();
  const p=personaSel.value;
  return DATA.filter(row=>{
    if(p){const t=tagsOf(row);
      if(p.startsWith("=")?!t.includes(p.slice(1)):!(t.includes(p)||t.includes("All")))return false;}
    if(g && !row.some(v=>String(v==null?"":v).toLowerCase().includes(g)))return false;
    for(let i=0;i<COLS.length;i++){
      if(colFilters[i] && !String(row[i]==null?"":row[i]).toLowerCase().includes(colFilters[i]))return false;
    }
    return true;
  });
}
function view(){
  let rows=filtered();
  if(sortCol>=0){rows=rows.slice().sort((a,b)=>{
    let x=a[sortCol],y=b[sortCol];
    const nx=parseFloat(x),ny=parseFloat(y);
    if(!isNaN(nx)&&!isNaN(ny)){return (nx-ny)*sortDir;}
    return String(x==null?"":x).localeCompare(String(y==null?"":y))*sortDir;
  });}
  return rows;
}
let viewRows=[];
function render(){
  const rows=viewRows=view();
  tbody.innerHTML="";
  const frag=document.createDocumentFragment();
  for(const row of rows){const tr=document.createElement("tr");
    const n=row[NAME];
    tr.className=(n in pending?"pending ":"")+(selected.has(n)?"selected":"");
    if(editing){const td=document.createElement("td");td.className="sel";
      const cb=document.createElement("input");cb.type="checkbox";cb.checked=selected.has(n);
      td.appendChild(cb);td.onclick=e=>{e.stopPropagation();toggleSel(row,e.shiftKey,e.target===cb?cb.checked:!selected.has(n));};
      tr.appendChild(td);}
    row.forEach((v,i)=>{const td=document.createElement("td");
      const s=v==null?"":String(v);td.textContent=s;td.title=s;
      if(s==="true")td.className="bool-true";else if(s==="false")td.className="bool-false";
      if(editing){
        if(i===SUP){td.classList.add("tog");td.title="Click to toggle";
          td.onclick=e=>{e.stopPropagation();setState(row,{sup:!row[SUP]});commit([row]);};}
        else if(i===SR||i===NOTE){td.classList.add("ed");td.title=(s?s+" \u2014 ":"")+"double-click to edit";
          td.onclick=e=>e.stopPropagation();td.ondblclick=e=>{e.stopPropagation();inlineEdit(td,row,i);};}
        else td.onclick=()=>openEd(row);
      }
      tr.appendChild(td);});
    frag.appendChild(tr);}
  tbody.appendChild(frag);
  count.textContent=rows.length+(rows.length!==DATA.length?" / "+DATA.length:"");
  const all=$("selAll");if(all){const k=rows.filter(r=>selected.has(r[NAME])).length;
    all.checked=k>0&&k===rows.length;all.indeterminate=k>0&&k<rows.length;
    all.onchange=()=>{rows.forEach(r=>all.checked?selected.add(r[NAME]):selected.delete(r[NAME]));render();};}
  updateSel();
  thead.querySelectorAll("th .arrow").forEach((el,i)=>el.textContent=(i===sortCol)?(sortDir>0?"▲":"▼"):"");
}
q.oninput=render;
personaSel.onchange=render;
document.getElementById("cols").onclick=()=>{showFilters=!showFilters;buildHead();render();};
document.getElementById("reset").onclick=()=>{q.value="";personaSel.value="";colFilters.fill("");sortCol=-1;buildHead();render();};
// --- Export (current filtered + sorted view) ---
function download(name,blob){const a=document.createElement("a");a.href=URL.createObjectURL(blob);
  a.download=name;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(a.href),1000);}
function exportCsv(){
  const esc=v=>{const s=v==null?"":String(v);return /[",\r\n]/.test(s)?'"'+s.replace(/"/g,'""')+'"':s;};
  const lines=[COLS,...view()].map(r=>r.map(esc).join(","));
  download("data_stream_inventory.csv",new Blob(["\uFEFF"+lines.join("\r\n")],{type:"text/csv;charset=utf-8"}));
}
// Minimal single-sheet .xlsx: SpreadsheetML parts in an uncompressed (stored) zip.
const CRC=(()=>{const t=new Uint32Array(256);for(let n=0;n<256;n++){let c=n;for(let k=0;k<8;k++)c=c&1?0xEDB88320^(c>>>1):c>>>1;t[n]=c>>>0;}return t;})();
function crc32(b){let c=0xFFFFFFFF;for(let i=0;i<b.length;i++)c=CRC[(c^b[i])&255]^(c>>>8);return (c^0xFFFFFFFF)>>>0;}
function zip(files){
  const enc=new TextEncoder(),parts=[],cdir=[];let off=0;
  for(const [name,text] of files){
    const n=enc.encode(name),d=enc.encode(text),crc=crc32(d);
    const h=new DataView(new ArrayBuffer(30));
    h.setUint32(0,0x04034b50,true);h.setUint16(4,20,true);h.setUint16(6,0x0800,true);
    h.setUint32(14,crc,true);h.setUint32(18,d.length,true);h.setUint32(22,d.length,true);h.setUint16(26,n.length,true);
    parts.push(h,n,d);
    const c=new DataView(new ArrayBuffer(46));
    c.setUint32(0,0x02014b50,true);c.setUint16(4,20,true);c.setUint16(6,20,true);c.setUint16(8,0x0800,true);
    c.setUint32(16,crc,true);c.setUint32(20,d.length,true);c.setUint32(24,d.length,true);c.setUint16(28,n.length,true);
    c.setUint32(42,off,true);
    cdir.push(c,n);off+=30+n.length+d.length;
  }
  const size=cdir.reduce((s,x)=>s+x.byteLength,0),e=new DataView(new ArrayBuffer(22));
  e.setUint32(0,0x06054b50,true);e.setUint16(8,files.length,true);e.setUint16(10,files.length,true);
  e.setUint32(12,size,true);e.setUint32(16,off,true);
  return new Blob([...parts,...cdir,e],{type:"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"});
}
function exportXlsx(){
  const x=s=>String(s).replace(/[\x00-\x08\x0B\x0C\x0E-\x1F]/g,"").replace(/&/g,"&amp;").replace(/</g,"&lt;").replace(/>/g,"&gt;");
  const col=i=>{let s="";for(i++;i>0;i=Math.floor((i-1)/26))s=String.fromCharCode(65+(i-1)%26)+s;return s;};
  const cell=(v,r,i)=>{const ref=col(i)+r;
    if(v==null||v==="")return "";
    if(typeof v==="number"&&isFinite(v))return `<c r="${ref}"><v>${v}</v></c>`;
    if(typeof v==="boolean")return `<c r="${ref}" t="b"><v>${v?1:0}</v></c>`;
    return `<c r="${ref}" t="inlineStr"${r===1?' s="1"':""}><is><t xml:space="preserve">${x(v)}</t></is></c>`;};
  const rows=[COLS,...view()];
  const body=rows.map((row,ri)=>`<row r="${ri+1}">${row.map((v,i)=>cell(v,ri+1,i)).join("")}</row>`).join("");
  const last=col(COLS.length-1)+rows.length;
  const sheet=`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">`+
    `<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>`+
    `<sheetData>${body}</sheetData><autoFilter ref="A1:${last}"/></worksheet>`;
  download("data_stream_inventory.xlsx",zip([
    ["[Content_Types].xml",`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>`],
    ["_rels/.rels",`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>`],
    ["xl/workbook.xml",`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Inventory" sheetId="1" r:id="rId1"/></sheets><definedNames><definedName name="_xlnm._FilterDatabase" localSheetId="0" hidden="1">Inventory!$A$1:$${col(COLS.length-1)}$${rows.length}</definedName></definedNames></workbook>`],
    ["xl/_rels/workbook.xml.rels",`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>`],
    ["xl/styles.xml",`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="2"><font><sz val="11"/><name val="Calibri"/></font><font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF1F4E78"/></patternFill></fill></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf/></cellStyleXfs><cellXfs count="2"><xf xfId="0"/><xf xfId="0" fontId="1" fillId="2" applyFont="1" applyFill="1"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>`],
    ["xl/worksheets/sheet1.xml",sheet],
  ]));
}
document.getElementById("csv").onclick=exportCsv;
document.getElementById("xlsx").onclick=exportXlsx;
// --- Editor: single-row dialog, multi-select bulk edit, inline cells ---
let editing=false,cur=null,lastSel=-1,byAsked=false;
const selected=new Set();
const ed=document.getElementById("ed"),bk=document.getElementById("bk"),$=id=>document.getElementById(id);
const REASONS=["Test","Internal","Initiative-specific","Duplicate","Deprecated","Other"];
document.querySelectorAll("select.reasons").forEach(el=>el.innerHTML=REASONS.map(r=>`<option>${r}</option>`).join(""));
$("edit").onclick=()=>{editing=!editing;document.body.classList.toggle("editing",editing);
  $("edit").textContent=editing?"Editing — click a row":"Edit classifications";buildHead();render();};
function getBy(){let b=localStorage.getItem("dcinv-user");
  if(!b&&!byAsked){byAsked=true;b=(prompt("Your name (recorded on corrections):")||"").trim();if(b)localStorage.setItem("dcinv-user",b);}
  return b||"";}
function setEntry(n,entry){if(same(entry,CORR.streams[n]))delete pending[n];else pending[n]=entry;}
// Record the desired state of one row as the minimal correction vs. the rules.
// Unspecified fields keep the row's current value.
function setState(row,{tags,sup,reason,note}){
  const n=row[NAME],r=RULE[n];
  tags=tags??tagsOf(row);sup=sup??!!row[SUP];reason=reason??String(row[SR]||"");
  note=(note??String(row[NOTE]||"")).trim();
  if(sup&&!reason)reason="Other";
  const e={};
  if(tags.length&&!sameTags(tags,r.p.split("; ")))e.persona=tags;
  const rr=ruleReasonFor(r,e.persona||r.p.split("; "));
  if(sup!==!!rr||(sup&&reason!==rr)){e.suppressed=sup;if(sup)e.reason=reason;}
  if(note)e.note=note;
  let entry=null;
  if(Object.keys(e).length){const prev=corrFor(n);
    entry=same(strip(prev),e)?prev:{...e,...(getBy()?{by:getBy()}:{}),at:new Date().toISOString().slice(0,10)};}
  setEntry(n,entry);
}
function commit(rows){localStorage.setItem(LSK,JSON.stringify(pending));rows.forEach(applyRow);render();updatePending();}
const reasonParts=s=>{const [rs,...rx]=String(s||"").split(": ");
  return REASONS.includes(rs)?[rs,rx.join(": ")]:["Other",String(s||"")];};
function openEd(row){
  cur=row;const n=row[NAME],r=RULE[n],c=corrFor(n)||{};
  $("edName").textContent=n;
  $("edRule").textContent="Rules say: "+r.p+" ("+r.b+")"+(r.s?" · suppressed: "+r.r:"");
  const tags=tagsOf(row);
  $("edP").innerHTML=TAGS.map(t=>`<label><input type="checkbox" value="${t}"${tags.includes(t)?" checked":""}> ${t}</label>`).join("");
  $("edS").checked=!!row[SUP];
  [$("edR").value,$("edRx").value]=reasonParts(row[SR]);
  $("edN").value=c.note||"";
  ed.showModal();
}
ed.addEventListener("close",()=>{
  const v=ed.returnValue,n=cur&&cur[NAME];if(!n||v==="cancel"||!v)return;
  if(v==="revert")setEntry(n,null);
  else{const sup=$("edS").checked;
    setState(cur,{tags:[...$("edP").querySelectorAll("input:checked")].map(i=>i.value),sup,
      reason:sup?[$("edR").value,$("edRx").value.trim()].filter(Boolean).join(": "):"",note:$("edN").value});}
  commit([cur]);
});
// Selection (checkbox column; shift-click selects a range of the rows shown)
function toggleSel(row,shift,on){
  const idx=viewRows.indexOf(row);
  const span=shift&&lastSel>=0?viewRows.slice(Math.min(lastSel,idx),Math.max(lastSel,idx)+1):[row];
  span.forEach(r=>on?selected.add(r[NAME]):selected.delete(r[NAME]));
  lastSel=idx;render();
}
function updateSel(){$("selbar").hidden=!(editing&&selected.size);$("selN").textContent=selected.size;}
$("selClear").onclick=()=>{selected.clear();lastSel=-1;render();};
$("bulk").onclick=()=>{
  $("bkTitle").textContent="Bulk edit "+selected.size+" stream"+(selected.size>1?"s":"");
  $("bkGo").textContent="Apply to "+selected.size;
  $("bkP").innerHTML=TAGS.map(t=>`<label><input type="checkbox" value="${t}"> ${t}</label>`).join("");
  ["bkPm","bkSm","bkNm"].forEach(id=>$(id).value="");$("bkRx").value=$("bkN").value="";$("bkR").value="Initiative-specific";
  bk.showModal();
};
bk.addEventListener("close",()=>{
  const v=bk.returnValue;if(!v||v==="cancel")return;
  const rows=DATA.filter(r=>selected.has(r[NAME]));
  if(v==="revert"){if(!confirm("Revert "+rows.length+" streams to the rules?"))return;rows.forEach(r=>setEntry(r[NAME],null));}
  else{const pm=$("bkPm").value,sm=$("bkSm").value,nm=$("bkNm").value;
    const picked=[...$("bkP").querySelectorAll("input:checked")].map(i=>i.value);
    const reason=[$("bkR").value,$("bkRx").value.trim()].filter(Boolean).join(": ");
    for(const row of rows){const st={},t=tagsOf(row);
      if(pm==="set"&&picked.length)st.tags=picked;
      if(pm==="add")st.tags=[...t,...picked.filter(x=>!t.includes(x))];
      if(pm==="remove")st.tags=t.filter(x=>!picked.includes(x));
      if(sm==="on"){st.sup=true;st.reason=reason;}
      if(sm==="off"){st.sup=false;st.reason="";}
      if(nm==="set")st.note=$("bkN").value;
      if(nm==="clear")st.note="";
      setState(row,st);}}
  commit(rows);
});
// Inline cell editing: Enter/blur saves, Escape cancels
function inlineEdit(td,row,i){
  const before=String(row[i]||"");td.contentEditable="true";td.classList.add("editing");td.focus();
  document.getSelection().selectAllChildren(td);
  let cancel=false;
  td.onkeydown=e=>{if(e.key==="Enter"){e.preventDefault();td.blur();}if(e.key==="Escape"){cancel=true;td.blur();}};
  td.onblur=()=>{td.contentEditable="false";const val=td.textContent.trim();
    if(cancel||val===before){render();return;}
    if(i===SR)setState(row,val?{sup:true,reason:val}:{sup:false,reason:""});else setState(row,{note:val});
    commit([row]);};
}
function updatePending(){const k=Object.keys(pending).length;
  $("pend").textContent=k?k+" pending correction"+(k>1?"s":""):"";$("exp").hidden=$("disc").hidden=!k;}
$("disc").onclick=()=>{if(!confirm("Discard all pending corrections in this browser?"))return;
  pending={};localStorage.removeItem(LSK);DATA.forEach(applyRow);render();updatePending();};
$("exp").onclick=()=>{
  const streams={...CORR.streams};
  for(const [n,e] of Object.entries(pending)){if(e===null)delete streams[n];else streams[n]=e;}
  const sorted=Object.fromEntries(Object.keys(streams).sort().map(k=>[k,streams[k]]));
  const out={org:ORG,updated:new Date().toISOString(),streams:sorted};
  download(ORG+".json",new Blob([JSON.stringify(out,null,2)+"\n"],{type:"application/json"}));
  alert("Saved "+ORG+".json.\nCommit it to DataCloud_Inventory/corrections/ and regenerate; then click 'Discard pending'.");
};
updatePending();
buildHead();render();
</script></body></html>"""


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser(description="Generate a Data Cloud stream/DLO/DMO inventory.")
    ap.add_argument("--org", default="jdo-oe0sdd", help="sf org alias (default: jdo-oe0sdd)")
    ap.add_argument("--out", default="output", help="output directory (default: ./output)")
    ap.add_argument("--api-version", default="67.0", help="Salesforce API version (default: 67.0)")
    ap.add_argument("--workers", type=int, default=16, help="parallel mapping queries (default: 16)")
    ap.add_argument("--persona-rules", default=DEFAULT_RULES,
                    help="persona tagging rules JSON (default: persona_rules.json beside this script)")
    ap.add_argument("--corrections", metavar="PATH",
                    help="collaborator corrections JSON (default: corrections/<org>.json beside this script)")
    ap.add_argument("--from-snapshot", metavar="PATH",
                    help="skip the org fetch and rebuild from a raw_snapshot.json (fast re-tagging)")
    args = ap.parse_args()

    from datetime import datetime, timezone
    os.makedirs(args.out, exist_ok=True)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    snap_path = os.path.join(args.out, "raw_snapshot.json")
    if args.from_snapshot:
        print(f"Loading snapshot {args.from_snapshot}...", file=sys.stderr)
        with open(args.from_snapshot, encoding="utf-8") as f:
            snap = json.load(f)
        streams, dmos, dlos, maps = snap["streams"], snap["dmos"], snap["dlos"], snap["maps"]
        generated = snap.get("generated", generated)
    else:
        print(f"Auth via sf CLI for org '{args.org}'...", file=sys.stderr)
        instance, token = get_session(args.org)
        client = Client(instance, token, args.api_version)
        print(f"  {instance}", file=sys.stderr)

        print("Fetching data streams...", file=sys.stderr)
        streams = fetch_streams(client)
        print("Fetching DMOs...", file=sys.stderr)
        dmos = fetch_dmos(client)
        print(f"  {len(dmos)} DMOs", file=sys.stderr)
        print("Fetching DLOs (for orphan detection)...", file=sys.stderr)
        dlos = fetch_dlos(client)
        print(f"  {len(dlos)} DLOs", file=sys.stderr)
        print("Resolving DLO->DMO mappings (reverse, per DMO)...", file=sys.stderr)
        maps = fetch_mappings(client, dmos, args.workers)
        print(f"  {len(maps)} unique mappings", file=sys.stderr)
        with open(snap_path, "w", encoding="utf-8") as f:
            json.dump({"org": args.org, "generated": generated, "streams": streams,
                       "dmos": dmos, "dlos": dlos, "maps": maps}, f)

    stream_dmos = build_stream_dmos(streams, maps)
    tagger = PersonaTagger(args.persona_rules)
    corr_path = args.corrections or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "corrections", f"{args.org}.json")
    corrections = load_corrections(corr_path)
    rows, rule_values = build_rows(streams, stream_dmos, dmos, tagger, corrections["streams"])
    breakdowns = build_breakdowns(rows, streams, dlos, stream_dmos)
    connectors = build_connectors(streams)

    sheets = {
        "Inventory": rows,
        "Unmapped Streams": breakdowns["Unmapped Streams"],
        "Suppressed Streams": breakdowns["Suppressed Streams"],
        **{f"{p} Datasets": persona_rows(rows, p) for p in tagger.personas},
        "Needs Attention": breakdowns["Needs Attention"],
        "Orphaned DLOs": breakdowns["Orphaned DLOs"],
        "Summary by Persona": breakdowns["Summary by Persona"],
        "Summary by Connector": breakdowns["Summary by Connector"],
        "Summary by Category": breakdowns["Summary by Category"],
        "Summary by Source Object": breakdowns["Summary by Source Object"],
        "Connectors": connectors,
    }

    xlsx_path = os.path.join(args.out, "data_stream_inventory.xlsx")
    html_path = os.path.join(args.out, "data_stream_inventory.html")
    write_xlsx(xlsx_path, sheets)
    write_html(html_path, f"Data Cloud Inventory — {args.org}", rows, args.org, generated,
               tagger, corrections, rule_values)

    mapped = sum(1 for r in rows if r["Is Mapped"])
    attn = len(breakdowns["Needs Attention"])
    print(
        f"\nDone. {len(rows)} streams | {mapped} mapped | {len(rows) - mapped} unmapped | "
        f"{attn} need attention | {len(breakdowns['Orphaned DLOs'])} orphaned DLOs",
        file=sys.stderr,
    )
    print(f"  {xlsx_path}", file=sys.stderr)
    print(f"  {html_path}", file=sys.stderr)
    print(f"  {len(breakdowns['Suppressed Streams'])} suppressed (flagged, still listed)",
          file=sys.stderr)
    unclassified = [r["Data Stream Name"] for r in rows if r["Persona"] == "Unclassified"]
    if unclassified:
        print(f"  {len(unclassified)} unclassified (add rules/overrides): "
              + ", ".join(unclassified), file=sys.stderr)


if __name__ == "__main__":
    main()
